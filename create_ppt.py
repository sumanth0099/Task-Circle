import collections 
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()
    
    # Define colors
    PRIMARY_COLOR = RGBColor(41, 128, 185) # Blue
    TEXT_COLOR = RGBColor(44, 62, 80) # Dark Gray
    
    # Helper to format titles
    def set_title(slide, title_text):
        title = slide.shapes.title
        title.text = title_text
        for paragraph in title.text_frame.paragraphs:
            paragraph.alignment = PP_ALIGN.LEFT
            paragraph.font.color.rgb = PRIMARY_COLOR
            paragraph.font.bold = True
    
    # 1. Title Slide
    slide_layout = prs.slide_layouts[0] # Title Layout
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "TaskCircle"
    title.text_frame.paragraphs[0].font.bold = True
    title.text_frame.paragraphs[0].font.color.rgb = PRIMARY_COLOR
    subtitle.text = "A Collaborative Task Management Application\n\n[Your Name / Team Name]\n[College / Organization Name]"
    
    # 2. Abstract
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    set_title(slide, "Abstract")
    content = slide.placeholders[1]
    content.text = "• TaskCircle centralizes project management by organizing work into distinct 'Circles'.\n" \
                   "• Provides secure Google OAuth authentication.\n" \
                   "• Integrates automated background scheduling for deadline notifications.\n" \
                   "• Features an AI assistant (Groq) to brainstorm and break down tasks.\n" \
                   "• Built on robust technologies: React 19, Node.js, PostgreSQL, and Redis."
    
    # 3. Introduction
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_title(slide, "Introduction")
    content = slide.placeholders[1]
    content.text = "• Why centralized task management?\n  - Modern teams need streamlined workflows.\n" \
                   "• What TaskCircle provides:\n  - Unified platform for task tracking and role management.\n" \
                   "• Concept of Circles:\n  - Organized group work with Public/Private settings.\n" \
                   "• Key Features:\n  - Real-time task tracking, automated notifications, AI-assisted productivity."

    # 4. Problem Statement
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_title(slide, "Problem Statement")
    content = slide.placeholders[1]
    content.text = "• Fragmented Tools: Teams scatter communication and tracking across apps.\n" \
                   "• Weak Role Management: Basic group apps lack structured hierarchies (Admin, Mod, Member).\n" \
                   "• Missed Deadlines: Lack of automated, customizable daily reminders.\n" \
                   "• Brainstorming Friction: Breaking large project goals into tasks manually is time-consuming."

    # 5. Objectives
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_title(slide, "Objectives")
    content = slide.placeholders[1]
    content.text = "• Centralize task management into secure, distinct 'Circles'.\n" \
                   "• Provide robust role-based access control.\n" \
                   "• Integrate an AI assistant for instant task breakdown and assistance.\n" \
                   "• Implement automated background scheduling for deadline notifications."

    # 6. Existing System
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_title(slide, "Existing System")
    content = slide.placeholders[1]
    content.text = "• Multiple Disconnected Apps: Separate tools for chat, AI, and task boards.\n" \
                   "• Manual Tracking: Requires manual effort to check and update due dates.\n" \
                   "• Weak Permissions: Simple group applications give all users equal access, leading to disorganized tasks.\n" \
                   "• No Native AI: Separate AI tools are disconnected from actual project data."

    # 7. Proposed System
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_title(slide, "Proposed System")
    content = slide.placeholders[1]
    content.text = "• Unified Management: Combines Circles, tasks, and AI in one platform.\n" \
                   "• Seamless Login: Secure, passwordless login via Google OAuth 2.0.\n" \
                   "• Automated Reminders: Cron jobs handle due-date and overdue notifications.\n" \
                   "• Built-in AI Assistant: Contextualizes requests within the workspace.\n" \
                   "• Robust Architecture: RESTful API powered by Node.js and PostgreSQL."

    # 8. System Architecture & File Structure
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_title(slide, "System Architecture & File Structure")
    content = slide.placeholders[1]
    content.text = "Architecture Flow:\n" \
                   "[Browser] -> [Nginx / React SPA] -> [Express API] -> [PostgreSQL + Redis]\n" \
                   "  ↳ Integrations: Google OAuth API, Groq AI API\n\n" \
                   "File Structure:\n" \
                   "frontend/src/ (components, pages, api, styles)\n" \
                   "backend/src/ (routes, services, middleware, db)\n" \
                   "backend/migrations/ (SQL schemas)"
    for paragraph in content.text_frame.paragraphs:
        paragraph.font.size = Pt(16)

    # 9. Data Flow
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_title(slide, "Data Flow")
    content = slide.placeholders[1]
    content.text = "Main Request Flow:\n" \
                   "User -> React UI -> API Client -> Express Router -> Auth Middleware -> Service Logic -> PostgreSQL -> React UI\n\n" \
                   "Background Flow:\n" \
                   "node-cron Scheduler -> PostgreSQL -> Generates Notifications\n\n" \
                   "Session Flow:\n" \
                   "Google OAuth -> Express -> Session stored in Redis -> Cookie sent to Client."
    for paragraph in content.text_frame.paragraphs:
        paragraph.font.size = Pt(16)

    # 10. Database Tables
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_title(slide, "Database Tables")
    content = slide.placeholders[1]
    content.text = "• users: id (PK), google_id, email, name\n" \
                   "• circles: id (PK), name, code, privacy, created_by (FK)\n" \
                   "• memberships: id (PK), circle_id (FK), user_id (FK), role, status\n" \
                   "• join_requests: id (PK), circle_id (FK), user_id (FK), status\n" \
                   "• tasks: id (PK), circle_id (FK), title, status, assigned_to (FK), due_date\n" \
                   "• notifications: id (PK), user_id (FK), type, message\n" \
                   "• notification_preferences: user_id (PK), alert preferences"
    for paragraph in content.text_frame.paragraphs:
        paragraph.font.size = Pt(14)

    # 11. Modules & Features
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_title(slide, "Modules & Features")
    content = slide.placeholders[1]
    content.text = "• Authentication: Google OAuth, secure session persistence.\n" \
                   "• Circle Management: Create/join via code, role management (Admin/Mod).\n" \
                   "• Task Management: Global 'My Tasks', status updates (TODO, IN_PROGRESS).\n" \
                   "• Notifications: Scheduled deadlines, real-time unread alerts.\n" \
                   "• AI Chat: Groq Llama 3.1 assistant for brainstorming."

    # 12. Technologies Used
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_title(slide, "Technologies Used")
    content = slide.placeholders[1]
    content.text = "• Frontend: React 19, Vite, React Router DOM, Vanilla CSS\n" \
                   "• Backend: Node.js, Express.js\n" \
                   "• Database & Cache: PostgreSQL, Redis\n" \
                   "• External APIs: Google OAuth 2.0, Groq API (Llama 3.1)\n" \
                   "• Key Libraries: node-cron, pg, Passport.js\n" \
                   "• Deployment: Docker, Docker Compose, Nginx, Render"

    # 13. Implementation & Working
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_title(slide, "Implementation & Working")
    content = slide.placeholders[1]
    content.text = "[ Placeholder for Implementation Workflows ]\n\n" \
                   "• Login flow\n• Circle creation/joining\n• Task creation/assignment\n" \
                   "• Notifications\n• AI assistant workflows"
    
    # 14. Results & Screenshots
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_title(slide, "Results & Screenshots")
    content = slide.placeholders[1]
    content.text = "[ Placeholder for Project Screenshots ]\n\n" \
                   "• Login Screen\n• Dashboard\n• Circle View\n" \
                   "• Task Management UI\n• AI Assistant Chat"

    # 15. Advantages
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_title(slide, "Advantages")
    content = slide.placeholders[1]
    content.text = "• Security: Passwordless Google OAuth authentication.\n" \
                   "• Organization: Role-based access control and unified tracking.\n" \
                   "• Automation: Background cron jobs reduce missed deadlines.\n" \
                   "• Intelligence: Built-in AI reduces planning overhead.\n" \
                   "• Scalability: Redis session management and Docker-based deployment."

    # 16. Future Scope
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_title(slide, "Future Scope")
    content = slide.placeholders[1]
    content.text = "• WebSockets: Implement Socket.io for real-time live UI updates.\n" \
                   "• Mobile App: Develop a React Native application for mobile access.\n" \
                   "• Advanced AI: Generate tasks automatically from meeting transcripts.\n" \
                   "• Integrations: Sync with Google Calendar / Outlook."

    # 17. Conclusion
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_title(slide, "Conclusion")
    content = slide.placeholders[1]
    content.text = "• TaskCircle successfully centralizes project management and collaboration.\n" \
                   "• It eliminates reliance on fragmented tools through organized 'Circles'.\n" \
                   "• Automates critical deadline tracking via background notifications.\n" \
                   "• Employs a robust, scalable architecture (React, Node, Postgres, Redis).\n" \
                   "• Demonstrates effective AI integration to boost team productivity."

    # 18. Thank You
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Thank You"
    title.text_frame.paragraphs[0].font.bold = True
    title.text_frame.paragraphs[0].font.color.rgb = PRIMARY_COLOR
    subtitle.text = "Questions?\n\nTaskCircle: A Collaborative Task Management Application"
    
    prs.save('TaskCircle_Presentation.pptx')

if __name__ == '__main__':
    create_presentation()
